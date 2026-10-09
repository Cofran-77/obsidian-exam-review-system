# -*- coding: utf-8 -*-
"""
考研政治错题本 · 单一真源 (Single Source of Truth) 编译与 Word 原生排版引擎
========================================================================
功能：
1. 自动扫描 `01-马原/02-错题/*.md` 及后续科目错题卡片（唯一真源）
2. 动态编译聚合根目录 `错题本.md`，保持与原子卡片 100% 同步
3. 生成专业级 Word 打印讲义 `错题本.docx`（原生样式、零 Markdown 杂质字符、彩色避坑框、诊断总表）
4. 可选：将生成文件复制到用户指定的云端硬盘目录
"""

import os
import sys
import glob
import re
import shutil
from datetime import date, datetime
from pathlib import Path

try:
    import yaml
except ImportError:
    yaml = None

try:
    import docx
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import qn, nsdecls
except ImportError:
    docx = None

# 确保在 Windows 控制台下输出 UTF-8
if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

WORKSPACE_DIR = os.path.dirname(os.path.abspath(__file__))
GOOGLE_DRIVE_DIR = os.environ.get('WRONG_ANSWER_GOOGLE_DRIVE_DIR')
VALID_STATUSES = {'待复习', '待攻关', '复习中', '已掌握', '已归档'}
VALID_REVIEW_RESULTS = {'未复习', 'again', 'hard', 'good'}
SUBJECT_ORDER = ['马原', '毛中特', '史纲', '德法', '当代']
SUBJECT_LABELS = {
    '马原': '第一部分：马克思主义基本原理（马原）',
    '毛中特': '第二部分：毛泽东思想和中国特色社会主义理论体系概论（毛中特·习思想）',
    '史纲': '第三部分：中国近现代史纲要（史纲）',
    '德法': '第四部分：思想道德与法治（德法）',
    '当代': '第五部分：当代世界经济与政治（当代·时政）',
}
SUBJECT_CANONICAL = {'近现代史': '史纲', '思修': '德法', '新思想': '毛中特'}
MODULE_ORDER = {
    '马原': [
        '马克思主义导论与唯物论',
        '唯物辩证法',
        '认识论',
        '唯物史观（历史唯物主义）',
        '马克思主义政治经济学',
        '科学社会主义',
    ],
    '毛中特': [
        '毛泽东思想',
        '邓小平理论与中特理论体系',
        '习思想总论与核心要义',
        '“五位一体”总体布局',
        '“四个全面”与战略保障',
    ],
}
DEFAULT_CARD_VALUES = {
    'status': '待复习',
    'review_result': '未复习',
    'review_count': 0,
    'question_type': '单选',
    'source_exam': '题库未标注',
    'source_page': None,
    'error_category': '',
    'error_subtype': '',
}

# ----------------------------------------------------------------------
# 1. 原子卡片解析与数据结构提取
# ----------------------------------------------------------------------

def parse_yaml_frontmatter(text):
    """解析 Markdown 文件开头的 YAML frontmatter"""
    fm_match = re.match(r'^---\s*\n(.*?)\n---\s*\n(.*)$', text, re.S)
    if not fm_match:
        return {}, text
    fm_text, body = fm_match.group(1), fm_match.group(2)
    if yaml is not None:
        meta = yaml.safe_load(fm_text) or {}
        if not isinstance(meta, dict):
            raise ValueError('frontmatter 必须解析为对象')
    else:
        meta = _parse_simple_yaml(fm_text)
    return meta, body


def _parse_simple_yaml(text):
    """Fallback parser for the repository's flat YAML when PyYAML is unavailable."""
    meta = {}
    current_list = None
    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        if not line.strip() or line.lstrip().startswith('#'):
            continue
        stripped = line.strip()
        if stripped.startswith('- ') and current_list is not None:
            current_list.append(_parse_scalar(stripped[2:].strip()))
            continue
        if ':' not in stripped:
            continue
        key, value = stripped.split(':', 1)
        key = key.strip()
        value = value.strip()
        if not value:
            meta[key] = []
            current_list = meta[key]
        else:
            meta[key] = _parse_scalar(value)
            current_list = None
    return meta


def _parse_scalar(value):
    value = value.strip()
    if len(value) >= 2 and value[0] == value[-1] and value[0] in "'\"":
        return value[1:-1]
    if value.lower() in {'null', '~'}:
        return None
    if value.lower() == 'true':
        return True
    if value.lower() == 'false':
        return False
    if re.fullmatch(r'-?\d+', value):
        return int(value)
    return value


def normalize_metadata(meta):
    normalized = dict(DEFAULT_CARD_VALUES)
    normalized.update(meta or {})
    normalized['review_count'] = int(normalized.get('review_count') or 0)
    if not normalized.get('error_category') and normalized.get('error_type'):
        normalized['error_category'] = str(normalized['error_type']).split('/', 1)[0]
    if not normalized.get('error_subtype') and normalized.get('error_type'):
        parts = str(normalized['error_type']).split('/', 1)
        normalized['error_subtype'] = parts[1] if len(parts) == 2 else parts[0]
    return normalized


def discover_question_files(workspace_dir=WORKSPACE_DIR):
    root = Path(workspace_dir)
    return sorted(
        path for path in root.rglob('*.md')
        if path.parent.name == '02-错题' and '.obsidian' not in path.parts
    )


def _valid_iso_date(value):
    if value in (None, ''):
        return True
    try:
        date.fromisoformat(str(value))
        return True
    except (TypeError, ValueError):
        return False


def validate_cards(cards, workspace_dir=WORKSPACE_DIR):
    errors = []
    seen = {}
    for index, card in enumerate(cards, start=1):
        card_id = card.get('id')
        if not card_id:
            errors.append(f'第 {index} 张卡片缺少 id')
        elif card_id in seen:
            errors.append(f'重复 id: {card_id}（第 {seen[card_id]} 张与第 {index} 张）')
        else:
            seen[card_id] = index
        if card.get('status') not in VALID_STATUSES:
            errors.append(f'{card_id or index} 的 status 不合法: {card.get("status")}')
        if card.get('review_result') not in VALID_REVIEW_RESULTS:
            errors.append(f'{card_id or index} 的 review_result 不合法: {card.get("review_result")}')
        for field in ('last_review', 'next_review'):
            if not _valid_iso_date(card.get(field)):
                errors.append(f'{card_id or index} 的 {field} 日期格式不合法: {card.get(field)}')
        for field in ('target_kp_id', 'thematic_hub_id'):
            value = card.get(field)
            if value and not re.fullmatch(r'[A-Z]+-[A-Z]+(?:-[A-Z0-9]+)+', str(value)):
                errors.append(f'{card_id or index} 的 {field} ID 格式不合法: {value}')
    return errors


def is_due(card, today=None):
    if card.get('status') == '已掌握':
        return False
    if not card.get('next_review'):
        return True
    today = today or date.today()
    try:
        return date.fromisoformat(str(card['next_review'])) <= today
    except ValueError:
        return False

def parse_question_card(file_path):
    """解析单个错题卡片，提取元数据与内容块"""
    with open(file_path, 'r', encoding='utf-8') as f:
        content = f.read()

    meta, body = parse_yaml_frontmatter(content)
    meta = normalize_metadata(meta)

    # 提取题目标题
    title_match = re.search(r'^#\s+(.+)$', body, re.M)
    raw_title = title_match.group(1) if title_match else os.path.splitext(os.path.basename(file_path))[0]
    # 清除前缀 Emoji 如 ❌
    clean_title = re.sub(r'^[❌📌⚠️\s]+', '', raw_title).strip()
    q_summary = re.sub(r'^错题\d+[·\s]+(?:\d+题[：:])?', '', clean_title).strip() or meta.get('topic', '核心考点')

    # 提取【题目原题与选项】
    q_sec = re.search(r'## 📝 题目原题与选项\s*\n(.*?)(?=\n##|\Z)', body, re.S)
    q_raw = q_sec.group(1).strip() if q_sec else ""
    q_lines = []
    ans_line = ""
    for line in q_raw.split('\n'):
        line = line.strip()
        if not line:
            continue
        if '【标准参考答案】' in line or '【参考答案】' in line:
            ans_line = re.sub(r'[\*\`]', '', line.split('：')[-1]).strip()
        else:
            q_lines.append(line)

    # 提取【选项深度剖析与干扰排除】
    ana_sec = re.search(r'## 🔍 选项深度剖析与干扰排除\s*\n(.*?)(?=\n##|\Z)', body, re.S)
    ana_raw = ana_sec.group(1).strip() if ana_sec else ""
    ana_items = []
    for line in ana_raw.split('\n'):
        line = line.strip()
        if not line:
            continue
        if line.startswith('- '):
            ana_items.append(line[2:].strip())
        elif re.match(r'^\d+\.\s+', line):
            ana_items.append(re.sub(r'^\d+\.\s+', '', line).strip())
        else:
            ana_items.append(line)

    # 提取【错因透视与解题复盘】(Callout 提示)
    tip_sec = re.search(r'## 💡 错因透视与解题复盘\s*\n(.*?)(?=\n##|\Z)', body, re.S)
    tip_raw = tip_sec.group(1).strip() if tip_sec else ""
    tip_text_lines = []
    for line in tip_raw.split('\n'):
        line = line.strip()
        if not line or line.startswith('>[!CAUTION]') or line.startswith('> [!CAUTION]'):
            continue
        if line.startswith('>'):
            clean_l = re.sub(r'^>\s*', '', line).strip()
            clean_l = clean_l.replace('**', '').replace('`', '')
            if clean_l:
                tip_text_lines.append(clean_l)
        else:
            clean_l = line.replace('**', '').replace('`', '')
            if clean_l:
                tip_text_lines.append(clean_l)
    tip_content = '\n'.join(tip_text_lines).strip()

    # 提取【关联知识点】
    kp_sec = re.search(r'## 🎯 错题对应具体知识点\s*\n(.*?)(?=\n##|\Z)', body, re.S)
    kp_raw = kp_sec.group(1).strip() if kp_sec else ""
    kp_links = re.findall(r'\[\[(.*?)\]\]', kp_raw)
    kp_name = kp_links[0] if kp_links else meta.get('target_kp', '').replace('[[', '').replace(']]', '')

    card_name = os.path.splitext(os.path.basename(file_path))[0]

    return {
        'file_path': file_path,
        'card_name': card_name,
        'meta': meta,
        'title': clean_title,
        'q_summary': q_summary,
        'q_num': meta.get('question_num', 999),
        'subject': meta.get('subject', '马原'),
        'module': meta.get('module', '未分类板块'),
        'topic': meta.get('topic', '核心考点'),
        'correct_ans': str(meta.get('correct_ans', ans_line or '待核对')),
        'user_choice': str(meta.get('user_choice', '未记录')),
        'error_type': meta.get('error_type') or '/'.join(filter(None, [meta.get('error_category'), meta.get('error_subtype')])) or '未分类偏差',
        'status': meta.get('status', '待复习'),
        'q_lines': q_lines,
        'ana_items': ana_items,
        'tip_content': tip_content,
        'kp_name': kp_name
    }

def load_all_questions():
    """扫描所有错题卡片并按题号排序"""
    files = [str(path) for path in discover_question_files()]
    questions = []
    cards = []
    for f in files:
        try:
            with open(f, 'r', encoding='utf-8') as source:
                meta, _ = parse_yaml_frontmatter(source.read())
            cards.append(normalize_metadata(meta))
            q = parse_question_card(f)
            questions.append(q)
        except Exception as e:
            print(f"Error parsing {f}: {e}")
    errors = validate_cards(cards)
    if errors:
        raise ValueError('错题卡片校验失败:\n- ' + '\n- '.join(errors))
    def get_sort_key(q):
        subj = SUBJECT_CANONICAL.get(q['subject'], q['subject'])
        subj_idx = SUBJECT_ORDER.index(subj) if subj in SUBJECT_ORDER else 99
        mod_order = MODULE_ORDER.get(subj, [])
        mod_idx = mod_order.index(q['module']) if q['module'] in mod_order else 99
        return (subj_idx, mod_idx, q['q_num'])
    questions.sort(key=get_sort_key)
    return questions

# ----------------------------------------------------------------------
# 2. 动态生成单一真源 错题本.md
# ----------------------------------------------------------------------

def compile_master_markdown(questions):
    """根据原子卡片编译根目录 错题本.md（单一真源聚合，内嵌内部跳转目录）"""
    # 统计各科目数量与按学科/板块归类
    subject_counts = {}
    subjects = {}
    for q in questions:
        subject = SUBJECT_CANONICAL.get(q['subject'], q['subject'])
        subjects.setdefault(subject, {}).setdefault(q['module'], []).append(q)
        subject_counts[subject] = subject_counts.get(subject, 0) + 1

    my_count = subject_counts.get('马原', 0)
    mzt_count = subject_counts.get('毛中特', 0)

    md_lines = [
        "# 📖 考研政治·错题精炼与核心考点复盘本（全局主库）\n",
        "> **架构说明**：本文件由各学科 `02-错题` 原子卡片自动编译生成，为只读汇总与集中复盘视图。",
        "> 错题录入或修改请直接编辑对应原子卡片，运行 `python format_docx.py` 自动同步本文档与 Word 排版讲义。\n",
        f"- **总收录错题数**：`{len(questions)}` 道（马原 {my_count} 题 · 毛中特 {mzt_count} 题）",
        "- **复习看板直达**：[[01-今日待复习队列|今日待复习]] ｜ [[02-认知偏差TOP分析榜|认知偏差 TOP 榜]] ｜ [[03-板块错题全景看板|板块错题看板]] ｜ [[04-核心知识点母库|核心知识点母库]] ｜ [[05-复习健康度|复习健康度]]\n",
        "## 📑 快速目录索引\n",
        "> **💡 交互提示**：点击下方各板块中的具体错题，可直接在本文档内平滑跳转至对应题目及深度解析；点击每道题目底部的 **【🔝 返回目录】** 随时返回顶端索引。\n",
    ]

    for subject in SUBJECT_ORDER:
        modules = subjects.get(subject, {})
        subj_total = subject_counts.get(subject, 0)
        label = SUBJECT_LABELS.get(subject, subject)
        md_lines.append(f"### 🏛️ {label}（共 {subj_total} 题）\n")

        if not modules:
            md_lines.append("> *（当前科目暂无错题录入，后续错题卡录入后自动挂载）*\n")
            continue

        order = MODULE_ORDER.get(subject, [])
        sorted_m_names = sorted(modules.keys(), key=lambda m: order.index(m) if m in order else 999)

        for m_name in sorted_m_names:
            m_qs = modules[m_name]
            m_qs.sort(key=lambda x: x['q_num'])
            md_lines.append(f"#### 📌 板块：{m_name}（共 {len(m_qs)} 题）")
            for q in m_qs:
                link_text = f"单选第{q['q_num']}题：{q['q_summary']}"
                md_lines.append(f"- [[#{link_text}|{link_text}]]")
            md_lines.append("")
        md_lines.append("")

    md_lines.extend(["\n---\n", "## 📝 错题正文\n"])

    for subject in SUBJECT_ORDER:
        modules = subjects.get(subject, {})
        if not modules:
            continue
        md_lines.append(f"# 🏛️ {SUBJECT_LABELS[subject]}\n")

        order = MODULE_ORDER.get(subject, [])
        sorted_m_names = sorted(modules.keys(), key=lambda m: order.index(m) if m in order else 999)

        for m_name in sorted_m_names:
            m_qs = modules[m_name]
            m_qs.sort(key=lambda x: x['q_num'])
            md_lines.append(f"## 📌 板块：{m_name}\n")
            md_lines.append(f"> **板块主页直达**：[[板块·{m_name}]]\n")

            for q in m_qs:
                q_heading = f"单选第{q['q_num']}题：{q['q_summary']}"
                md_lines.append(f"### {q_heading}\n")
                md_lines.append(f"> **📄 独立原子卡片**：[[{q['card_name']}]]  ")
                md_lines.append(f"> **做题对决**：标准答案 `【{q['correct_ans']}】` 🆚 你的选择 `【{q['user_choice']}】`  ")
                md_lines.append(f"> **受控认知偏差**：`{q['error_type']}`  ")
                md_lines.append(f"> **考点归属**：{q['topic']}  ")
                if q['kp_name']:
                    md_lines.append(f"> **母库知识点**：[[{q['kp_name']}]]\n")
                else:
                    md_lines.append("\n")

                md_lines.append("#### 📝 题目原题与选项\n")
                for ql in q['q_lines']:
                    md_lines.append(f"{ql}  ")
                md_lines.append(f"\n- **【标准参考答案】**：`{q['correct_ans']}`\n")

                if q['ana_items']:
                    md_lines.append("#### 🔍 选项深度剖析与干扰排除\n")
                    for item in q['ana_items']:
                        md_lines.append(f"- {item}")
                    md_lines.append("")

                if q['tip_content']:
                    md_lines.append("#### 💡 错因透视与避坑速记\n")
                    md_lines.append("> [!CAUTION]")
                    for tip_l in q['tip_content'].split('\n'):
                        md_lines.append(f"> {tip_l}")
                    md_lines.append("")

                md_lines.append(f"👉 [[#📑 快速目录索引|🔝 返回目录]] ｜ [[#📌 板块：{m_name}|⬆️ 返回所属板块]] ｜ [[{q['card_name']}|📄 打开独立卡片]]\n")
                md_lines.append("---\n")

    out_content = '\n'.join(md_lines)
    target_path = os.path.join(WORKSPACE_DIR, '错题本.md')
    with open(target_path, 'w', encoding='utf-8') as f:
        f.write(out_content)
    print(f"✓ Master Markdown successfully compiled: {target_path} ({len(questions)} questions)")

# ----------------------------------------------------------------------
# 3. 生成专业级 Word 错题本 (format_docx)
# ----------------------------------------------------------------------

def set_cell_background(cell, fill_hex):
    """设置单元格背景颜色"""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """设置单元格内边距"""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_formatted_text(paragraph, text, default_color=None, default_size=Pt(10.5)):
    """解析文本中的 **加粗** 并生成真正的 Word 加粗 run，彻底去除 ** 符号"""
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            content = part[2:-2]
            run = paragraph.add_run(content)
            run.bold = True
            run.font.name = '微软雅黑'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
            run.font.size = default_size
            if default_color:
                run.font.color.rgb = default_color
        else:
            clean_part = part.replace('`', '')
            run = paragraph.add_run(clean_part)
            run.font.name = '宋体'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
            run.font.size = default_size
            if default_color:
                run.font.color.rgb = default_color

def build_docx_handout(questions):
    """构建精美排版的 Word 讲义"""
    if docx is None:
        print("! 未安装 python-docx，已跳过 Word 导出；Markdown 汇总仍会生成。")
        return
    doc = docx.Document()

    # 页面边距 2cm (0.8英寸)
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    # 1. 封面主标题
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(14)
    p_title.paragraph_format.space_after = Pt(4)
    run_t = p_title.add_run('考研政治 · 错题精炼与核心考点复盘手册')
    run_t.font.name = '黑体'
    run_t._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    run_t.font.size = Pt(22)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(30, 41, 59)

    # 副标题与元数据
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run(f'基于《2027考研政治1000题》全量诊断 · 认知偏差避坑速记库（收录 {len(questions)} 题）')
    run_sub.font.name = '楷体'
    run_sub._element.rPr.rFonts.set(qn('w:eastAsia'), '楷体')
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = RGBColor(100, 116, 139)

    # 2. 全局诊断速览表格（错题体检单）
    p_diag_h = doc.add_paragraph()
    p_diag_h.paragraph_format.space_before = Pt(8)
    p_diag_h.paragraph_format.space_after = Pt(4)
    r_dh = p_diag_h.add_run('📋 错题全景诊断与对决一览表')
    r_dh.font.name = '微软雅黑'
    r_dh._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    r_dh.font.size = Pt(12)
    r_dh.font.bold = True
    r_dh.font.color.rgb = RGBColor(15, 23, 42)

    table = doc.add_table(rows=len(questions) + 1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(0.7), Inches(1.5), Inches(1.8), Inches(0.8), Inches(0.8), Inches(1.4)]
    headers = ['题号', '所属板块', '命题考点', '正答', '作答', '认知偏差归类']

    # 表头样式
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], "1E293B") # 深蓝灰底纹
        set_cell_margins(hdr_cells[i], top=80, bottom=80, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = '微软雅黑'
            r._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
            r.font.size = Pt(9)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    # 填充表格内容
    for idx, q in enumerate(questions):
        row_cells = table.rows[idx + 1].cells
        # 交替行底纹
        bg_color = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
        for c in row_cells:
            set_cell_background(c, bg_color)
            set_cell_margins(c, top=60, bottom=60, left=80, right=80)

        # 题号
        row_cells[0].paragraphs[0].text = f"[{q['subject']}] 题{q['q_num']}"
        row_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        # 板块
        row_cells[1].paragraphs[0].text = q['module'].replace('马克思主义', '')
        # 考点 / 命题概要
        row_cells[2].paragraphs[0].text = q['q_summary']
        # 正答
        p_ans = row_cells[3].paragraphs[0]
        p_ans.text = q['correct_ans']
        p_ans.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p_ans.runs:
            p_ans.runs[0].font.bold = True
            p_ans.runs[0].font.color.rgb = RGBColor(22, 101, 52) # 绿色
        # 作答
        p_usr = row_cells[4].paragraphs[0]
        p_usr.text = q['user_choice']
        p_usr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p_usr.runs:
            p_usr.runs[0].font.bold = True
            p_usr.runs[0].font.color.rgb = RGBColor(185, 28, 28) # 红色
        # 偏差
        p_err = row_cells[5].paragraphs[0]
        p_err.text = q['error_type']
        if p_err.runs:
            p_err.runs[0].font.size = Pt(8.5)
            p_err.runs[0].font.color.rgb = RGBColor(67, 56, 202) # 紫蓝

        # 统一字体大小
        for i in range(6):
            row_cells[i].width = col_widths[i]
            for r in row_cells[i].paragraphs[0].runs:
                r.font.name = '宋体'
                r._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
                if not r.font.size:
                    r.font.size = Pt(8.5)

    doc.add_page_break()

    # 3. 按科目与板块层次输出错题卡片
    subjects = {}
    for q in questions:
        subject = SUBJECT_CANONICAL.get(q['subject'], q['subject'])
        subjects.setdefault(subject, {}).setdefault(q['module'], []).append(q)

    for subject in SUBJECT_ORDER:
        modules = subjects.get(subject, {})
        if not modules:
            continue
        p_subj = doc.add_paragraph()
        p_subj.paragraph_format.space_before = Pt(22)
        p_subj.paragraph_format.space_after = Pt(8)
        r_subj = p_subj.add_run(f'🏛️ 【科目】{SUBJECT_LABELS[subject]}')
        r_subj.font.name = '黑体'
        r_subj._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        r_subj.font.size = Pt(16)
        r_subj.font.bold = True
        r_subj.font.color.rgb = RGBColor(15, 23, 42)

        order = MODULE_ORDER.get(subject, [])
        sorted_m_names = sorted(modules.keys(), key=lambda m: order.index(m) if m in order else 999)

        for m_name in sorted_m_names:
            m_qs = modules[m_name]
            m_qs.sort(key=lambda x: x['q_num'])

            # 板块标题
            p_sec = doc.add_paragraph()
            p_sec.paragraph_format.space_before = Pt(14)
            p_sec.paragraph_format.space_after = Pt(6)
            r_sec = p_sec.add_run(f'📑 【核心板块】{m_name}（共 {len(m_qs)} 题）')
            r_sec.font.name = '黑体'
            r_sec._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
            r_sec.font.size = Pt(13)
            r_sec.font.bold = True
            r_sec.font.color.rgb = RGBColor(30, 41, 59)

            for q in m_qs:
                # 题目标题
                p_q_title = doc.add_paragraph()
                p_q_title.paragraph_format.space_before = Pt(12)
                p_q_title.paragraph_format.space_after = Pt(4)
                r_qt = p_q_title.add_run(f'📌 [{q["subject"]}] 单选第 {q["q_num"]} 题：{q["q_summary"]}')
                r_qt.font.name = '微软雅黑'
                r_qt._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
                r_qt.font.size = Pt(11.5)
                r_qt.font.bold = True
                r_qt.font.color.rgb = RGBColor(30, 64, 175)

                # 对决记录与认知偏差栏
                p_bar = doc.add_paragraph()
                p_bar.paragraph_format.space_before = Pt(2)
                p_bar.paragraph_format.space_after = Pt(4)

                # 标准答案
                r1 = p_bar.add_run('【参考答案】')
                r1.font.bold = True
                r1.font.color.rgb = RGBColor(22, 101, 52)
                r1_v = p_bar.add_run(f' {q["correct_ans"]}   ')
                r1_v.font.bold = True
                r1_v.font.size = Pt(11)
                r1_v.font.color.rgb = RGBColor(22, 101, 52)

                # 你的作答
                r2 = p_bar.add_run('【你的作答】')
                r2.font.bold = True
                r2.font.color.rgb = RGBColor(185, 28, 28)
                r2_v = p_bar.add_run(f' {q["user_choice"]}   ')
                r2_v.font.bold = True
                r2_v.font.size = Pt(11)
                r2_v.font.color.rgb = RGBColor(185, 28, 28)

                # 认知偏差
                r3 = p_bar.add_run('【认知偏差】')
                r3.font.bold = True
                r3.font.color.rgb = RGBColor(67, 56, 202)
                r3_v = p_bar.add_run(f' {q["error_type"]}')
                r3_v.font.name = '微软雅黑'
                r3_v._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
                r3_v.font.color.rgb = RGBColor(67, 56, 202)

                for r in [r1, r2, r3]:
                    r.font.name = '微软雅黑'
                    r._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
                    r.font.size = Pt(10)

                # 题目原题与选项
                if q['q_lines']:
                    for ql in q['q_lines']:
                        p_ql = doc.add_paragraph()
                        p_ql.paragraph_format.space_after = Pt(2)
                        p_ql.paragraph_format.line_spacing = 1.2
                        add_formatted_text(p_ql, ql, default_size=Pt(10.5))

                # 官方精析与干扰排除
                if q['ana_items']:
                    p_ana_h = doc.add_paragraph()
                    p_ana_h.paragraph_format.space_before = Pt(4)
                    p_ana_h.paragraph_format.space_after = Pt(2)
                    r_ah = p_ana_h.add_run('【解析要点与干扰排除】')
                    r_ah.font.bold = True
                    r_ah.font.name = '微软雅黑'
                    r_ah._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
                    r_ah.font.size = Pt(10)
                    r_ah.font.color.rgb = RGBColor(30, 41, 59)

                    for item in q['ana_items']:
                        p_item = doc.add_paragraph(style='List Bullet')
                        p_item.paragraph_format.space_after = Pt(2)
                        p_item.paragraph_format.line_spacing = 1.15
                        add_formatted_text(p_item, item, default_size=Pt(9.5))

                # 避坑速记框 (单格优雅底纹表格)
                if q['tip_content']:
                    tip_tbl = doc.add_table(rows=1, cols=1)
                    tip_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    tip_tbl.autofit = False
                    cell = tip_tbl.cell(0, 0)
                    cell.width = Inches(6.8)
                    set_cell_background(cell, "FEF9C3") # 淡黄色底纹
                    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)

                    # 边框：仅左侧粗橙色边框
                    tcPr = cell._element.get_or_add_tcPr()
                    borders = parse_xml(
                        f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:left w:val="single" w:sz="18" w:space="0" w:color="F59E0B"/>'
                        f'<w:top w:val="none"/>'
                        f'<w:right w:val="none"/>'
                        f'<w:bottom w:val="none"/>'
                        f'</w:tcBorders>'
                    )
                    tcPr.append(borders)

                    p_box = cell.paragraphs[0]
                    p_box.paragraph_format.space_before = Pt(0)
                    p_box.paragraph_format.space_after = Pt(0)

                    r_bulb = p_box.add_run('💡 避坑速记 / 核心解题套路：\n')
                    r_bulb.font.bold = True
                    r_bulb.font.name = '微软雅黑'
                    r_bulb._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
                    r_bulb.font.size = Pt(9.5)
                    r_bulb.font.color.rgb = RGBColor(180, 83, 9)

                    r_tip = p_box.add_run(q['tip_content'])
                    r_tip.font.name = '楷体'
                    r_tip._element.rPr.rFonts.set(qn('w:eastAsia'), '楷体')
                    r_tip.font.size = Pt(9.5)
                    r_tip.font.color.rgb = RGBColor(120, 53, 15)

                # 题末分隔留白
                p_sep = doc.add_paragraph()
                p_sep.paragraph_format.space_before = Pt(6)

    out_docx_path = os.path.join(WORKSPACE_DIR, '错题本.docx')
    doc.save(out_docx_path)
    print(f"✓ Professional Word Document generated: {out_docx_path}")

# ----------------------------------------------------------------------
# 4. 云端硬盘同步 (Google Drive)
# ----------------------------------------------------------------------

def sync_to_google_drive():
    """将生成的 错题本.docx 与 错题本.md 同步到 Google Drive"""
    if not GOOGLE_DRIVE_DIR:
        return
    if not os.path.exists(GOOGLE_DRIVE_DIR):
        print(f"! Google Drive directory not found at: {GOOGLE_DRIVE_DIR}")
        return

    files_to_sync = ['错题本.docx', '错题本.md']
    for fname in files_to_sync:
        src = os.path.join(WORKSPACE_DIR, fname)
        dst = os.path.join(GOOGLE_DRIVE_DIR, fname)
        if os.path.exists(src):
            try:
                shutil.copy2(src, dst)
                print(f"✓ Synced to Google Drive: {dst}")
            except Exception as e:
                print(f"! Failed to sync {fname} to Google Drive: {e}")

# ----------------------------------------------------------------------
# 主执行入口
# ----------------------------------------------------------------------

def main():
    print("==================================================")
    print("🚀 启动考研政治错题本 SSOT 单一真源编译与排版流水线...")
    print("==================================================")
    try:
        questions = load_all_questions()
    except ValueError as exc:
        print(f"! {exc}")
        return 1
    print(f"1. 成功扫描并提取 {len(questions)} 张错题原子卡片元数据。")

    compile_master_markdown(questions)
    build_docx_handout(questions)
    sync_to_google_drive()

    print("==================================================")
    print("🎉 错题本编译与同步全部顺利完成！")
    print("==================================================")
    return 0

if __name__ == '__main__':
    raise SystemExit(main())
